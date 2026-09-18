from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).parent
OUT = ROOT / 'Orbit_App_Documentation.docx'
PURPLE = RGBColor(91, 92, 226)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(.7)
section.bottom_margin = Inches(.7)
section.left_margin = Inches(.75)
section.right_margin = Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
styles['Title'].font.name = 'Aptos Display'
styles['Title'].font.size = Pt(29)
styles['Title'].font.color.rgb = RGBColor(35, 35, 56)
styles['Heading 1'].font.name = 'Aptos Display'
styles['Heading 1'].font.size = Pt(16)
styles['Heading 1'].font.color.rgb = RGBColor(46, 46, 77)

def eyebrow(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(9); r.font.color.rgb = PURPLE

def title(text):
    doc.add_paragraph(text, style='Title')

def heading(text):
    doc.add_paragraph(text, style='Heading 1')

def paragraph(text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(7)
    return p

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

# Page 1
eyebrow('FLUTTER APPLICATION DOCUMENTATION')
title('Orbit Registration App')
p = paragraph('A polished, responsive three-screen Flutter experience that guides a user from introduction to account creation and a personalized confirmation.')
p.runs[0].font.size = Pt(13); p.runs[0].font.color.rgb = RGBColor(94, 94, 115)
heading('Purpose')
paragraph('Orbit is designed as a concise registration journey with a friendly editorial tone. The interface makes the next action obvious, gives useful feedback while data is entered, and confirms exactly what was submitted at the end of the flow.')
heading('Screen flow')
for label, text in [
    ('1. Home', 'Sets the product tone with a focused value proposition, two benefits, and a prominent account-creation action.'),
    ('2. Registration form', 'Collects identity and account details with simple controls, clear labels, and inline validation.'),
    ('3. Detail', 'Welcomes the user by name and shows the submitted email address and selected profile type.'),
    ('Named navigation', 'The route structure is / (Home) → /register (Registration) → /detail (Detail, with submitted profile data).'),
]:
    p = doc.add_paragraph(); p.add_run(label + ': ').bold = True; p.add_run(text)
heading('Design approach')
paragraph('The visual system uses Material 3 with a violet seed color, generous rounded corners, subtle lavender surfaces, and compact iconography. Content is constrained to a readable maximum width while horizontal padding adapts on smaller screens, giving the experience a comfortable mobile and desktop presentation.')
paragraph('Prepared for the Orbit Flutter multi-screen application.').runs[0].italic = True

# Page 2
doc.add_page_break()
eyebrow('IMPLEMENTATION & EXPERIENCE')
title('How the app works')
heading('Interactive registration')
paragraph('The form is stateful and uses Flutter’s Form and TextFormField validation pattern. Users receive targeted error messages before they can continue, reducing ambiguity and making correction easy.')
for text in [
    'Required name: prevents a blank submission.',
    'Email validation: checks that the input follows a standard email format.',
    'Password validation: requires a password and enforces a minimum of eight characters.',
    'Password visibility: the trailing icon lets users reveal or conceal their entry.',
    'Profile selection: Choice Chips let users identify as a Creator, Founder, or Student.',
    'Terms acknowledgement: an unchecked agreement is clearly marked as required.',
]: bullet(text)
heading('Responsive behavior')
paragraph('Each screen sits inside a reusable shell that caps the content width at 620 pixels. On narrow devices, the home screen reduces its side padding; the form uses a scrollable list so the keyboard and smaller viewport never hide critical controls. Full-width primary buttons preserve an easy touch target across screen sizes.')
heading('Data hand-off and completion')
paragraph('Once all validation passes, the app creates a profile object containing the user’s name, email, and chosen role. That object is passed to the Detail route as route arguments. The final screen turns the name into an initial avatar and gives the user a clean summary, with a Back to home action that resets the journey naturally.')
heading('Quality checks completed')
paragraph('The app has been checked with Flutter’s static analyzer and has a widget test that verifies the primary Home-to-Registration navigation. Both checks pass without issues.')
p = paragraph('Key UX outcome: the application is not a set of static screens—the choices, form errors, password control, route transitions, submitted data, and return action all respond to the user.')
p.runs[0].bold = True; p.runs[0].font.color.rgb = PURPLE

# Visual appendix
for file_name, heading_text, caption in [
    ('Screenshot 2026-09-13 at 11.51.21 PM.png', 'Home screen', 'Figure 1. The initial Orbit onboarding screen.'),
    ('Screenshot 2026-09-13 at 11.52.15 PM.png', 'Registration form', 'Figure 2. Responsive registration form with validation controls.'),
    ('Screenshot 2026-09-13 at 11.52.34 PM.png', 'Account detail', 'Figure 3. Personalized account confirmation screen.'),
]:
    doc.add_page_break()
    eyebrow('VISUAL APPENDIX')
    title(heading_text)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT / file_name), height=Inches(7.4))
    p = doc.add_paragraph(caption); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True; p.runs[0].font.size = Pt(9)

doc.core_properties.title = 'Orbit Registration App — Documentation'
doc.core_properties.subject = 'Flutter multi-screen application documentation'
doc.save(OUT)
print(OUT)
